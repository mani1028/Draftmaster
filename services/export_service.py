import io
import re
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet
from docx import Document

class ExportService:
    @staticmethod
    def to_pdf(text):
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=letter)
        styles = getSampleStyleSheet()
        story = []
        
        # Clean up common markdown that might leak in
        text = text.replace('**', '')
        
        # Split by newlines
        paragraphs = text.split('\n')
        for p in paragraphs:
            if p.strip():
                # ReportLab Paragraph supports <b> tags natively
                story.append(Paragraph(p, styles['Normal']))
                story.append(Spacer(1, 12))
        
        doc.build(story)
        buffer.seek(0)
        return buffer

    @staticmethod
    def to_docx(text):
        doc = Document()
        doc.add_heading('Generated Document', 0)
        
        # Clean up markdown stars
        text = text.replace('**', '')
        
        paragraphs = text.split('\n')
        for p_text in paragraphs:
            if p_text.strip():
                p = doc.add_paragraph()
                
                # Split the paragraph text by <b> and </b> tags using regex
                parts = re.split(r'(<b>.*?</b>)', p_text)
                
                for part in parts:
                    if part.startswith('<b>') and part.endswith('</b>'):
                        # Remove the tags and add as bold run
                        clean_part = part.replace('<b>', '').replace('</b>', '')
                        p.add_run(clean_part).bold = True
                    else:
                        # Add as normal run
                        p.add_run(part)
                
        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        return buffer